import os
import zipfile
import io
import fitz  # PyMuPDF
from PIL import Image
from docx import Document
import ebooklib
from ebooklib import epub


ALLOWED_EXTENSIONS = {"cbz", "pdf", "docx", "epub"}
COVER_SIZE = (300, 420)


def allowed_file(filename):
    return (
        "." in filename
        and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS
    )


def get_file_type(filename):
    return filename.rsplit(".", 1)[1].lower() if "." in filename else ""


# ---------------------------------------------------------------------------
# Metadata extraction
# ---------------------------------------------------------------------------

def extract_metadata(file_path, file_type):
    """Return dict with title, author, description, page_count."""
    meta = {"title": "", "author": "", "description": "", "page_count": 0}
    try:
        if file_type == "pdf":
            meta.update(_meta_pdf(file_path))
        elif file_type == "epub":
            meta.update(_meta_epub(file_path))
        elif file_type == "docx":
            meta.update(_meta_docx(file_path))
        elif file_type == "cbz":
            meta.update(_meta_cbz(file_path))
    except Exception:
        pass
    return meta


def _meta_pdf(path):
    doc = fitz.open(path)
    info = doc.metadata or {}
    return {
        "title": info.get("title", ""),
        "author": info.get("author", ""),
        "description": info.get("subject", ""),
        "page_count": doc.page_count,
    }


def _meta_epub(path):
    book = epub.read_epub(path)
    meta = {
        "title": "",
        "author": "",
        "description": "",
        "page_count": 0,
    }
    titles = book.get_metadata("DC", "title")
    if titles:
        meta["title"] = titles[0][0]
    creators = book.get_metadata("DC", "creator")
    if creators:
        meta["author"] = creators[0][0]
    descs = book.get_metadata("DC", "description")
    if descs:
        meta["description"] = descs[0][0]
    items = list(book.get_items_of_type(ebooklib.ITEM_DOCUMENT))
    meta["page_count"] = len(items)
    return meta


def _meta_docx(path):
    doc = Document(path)
    props = doc.core_properties
    return {
        "title": props.title or "",
        "author": props.author or "",
        "description": props.description or "",
        "page_count": len(doc.paragraphs),
    }


def _meta_cbz(path):
    with zipfile.ZipFile(path) as zf:
        images = sorted(
            f for f in zf.namelist() if _is_image(f) and not f.startswith("__")
        )
    return {"title": "", "author": "", "description": "", "page_count": len(images)}


# ---------------------------------------------------------------------------
# Cover extraction
# ---------------------------------------------------------------------------

def extract_cover(file_path, file_type, cover_dir, book_id):
    """Extract/generate a cover image and save it; return the saved path or ''."""
    os.makedirs(cover_dir, exist_ok=True)
    out_path = os.path.join(cover_dir, f"{book_id}.jpg")
    try:
        if file_type == "pdf":
            _cover_pdf(file_path, out_path)
        elif file_type == "epub":
            _cover_epub(file_path, out_path)
        elif file_type == "cbz":
            _cover_cbz(file_path, out_path)
        elif file_type == "docx":
            _cover_docx(file_path, out_path)
        else:
            return ""
        return out_path if os.path.exists(out_path) else ""
    except Exception:
        return ""


def _save_pil(img, out_path):
    img = img.convert("RGB")
    img.thumbnail(COVER_SIZE, Image.LANCZOS)
    img.save(out_path, "JPEG", quality=85)


def _cover_pdf(src, out):
    doc = fitz.open(src)
    page = doc[0]
    pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
    _save_pil(img, out)


def _cover_epub(src, out):
    book = epub.read_epub(src)
    # Try cover item
    for item in book.get_items():
        if item.get_type() == ebooklib.ITEM_COVER:
            img = Image.open(io.BytesIO(item.get_content()))
            _save_pil(img, out)
            return
    # Fallback: first image
    for item in book.get_items_of_type(ebooklib.ITEM_IMAGE):
        try:
            img = Image.open(io.BytesIO(item.get_content()))
            _save_pil(img, out)
            return
        except Exception:
            continue


def _cover_cbz(src, out):
    with zipfile.ZipFile(src) as zf:
        images = sorted(
            f for f in zf.namelist() if _is_image(f) and not f.startswith("__")
        )
        if not images:
            return
        data = zf.read(images[0])
    img = Image.open(io.BytesIO(data))
    _save_pil(img, out)


def _cover_docx(src, out):
    doc = Document(src)
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            img_part = rel.target_part
            img = Image.open(io.BytesIO(img_part.blob))
            _save_pil(img, out)
            return


def _is_image(name):
    return name.lower().endswith((".jpg", ".jpeg", ".png", ".gif", ".webp"))


# ---------------------------------------------------------------------------
# CBZ page serving
# ---------------------------------------------------------------------------

def cbz_page_count(file_path):
    with zipfile.ZipFile(file_path) as zf:
        images = sorted(
            f for f in zf.namelist() if _is_image(f) and not f.startswith("__")
        )
    return len(images)


def cbz_get_page(file_path, page_index):
    """Return (image_bytes, mimetype) for the given 0-based page index."""
    with zipfile.ZipFile(file_path) as zf:
        images = sorted(
            f for f in zf.namelist() if _is_image(f) and not f.startswith("__")
        )
        if page_index < 0 or page_index >= len(images):
            return None, None
        data = zf.read(images[page_index])
        ext = images[page_index].rsplit(".", 1)[-1].lower()
        mime = "image/jpeg" if ext in ("jpg", "jpeg") else f"image/{ext}"
    return data, mime


# ---------------------------------------------------------------------------
# DOCX → HTML
# ---------------------------------------------------------------------------

def docx_to_html(file_path):
    """Convert a DOCX file to a simple HTML string."""
    doc = Document(file_path)
    parts = ["<div class='docx-content'>"]
    for para in doc.paragraphs:
        style = para.style.name if para.style else ""
        text = para.text
        if style.startswith("Heading 1"):
            parts.append(f"<h1>{text}</h1>")
        elif style.startswith("Heading 2"):
            parts.append(f"<h2>{text}</h2>")
        elif style.startswith("Heading 3"):
            parts.append(f"<h3>{text}</h3>")
        elif text.strip():
            parts.append(f"<p>{text}</p>")
        else:
            parts.append("<br>")
    parts.append("</div>")
    return "\n".join(parts)
